import io
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots

# --- 1. CONFIGURACIÓN VISUAL GENERAL ---
st.set_page_config(
    page_title="Management Suite | GlobalTech BI",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

PALETA_COLORES = ["#FFB800", "#1E2229", "#FF7A00", "#6C5CE7", "#00B894", "#E17055"]

# --- 2. FUNCIONES DE LIMPIEZA Y FORMATEO ---
def clean_val(v):
    if pd.isna(v) or v is None:
        return 0.0
    s = str(v).strip().replace('$', '').replace(' ', '').replace('\xa0', '')
    if not s or s.lower() in ('nan', 'none', 'null'):
        return 0.0
    if '.' in s and ',' in s:
        if s.rfind('.') > s.rfind(','):
            s = s.replace(',', '')
        else:
            s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) in (1, 2):
            s = s.replace(',', '.')
        else:
            s = s.replace(',', '')
    try:
        return float(s)
    except Exception:
        return 0.0

def render_kpi_card(titulo, val_act, val_ant, val_meta, badge_txt, badge_class, fill_class, es_moneda=False):
    """Genera tarjeta KPI ejecutiva con formato de píldora y barra de progreso."""
    if es_moneda:
        txt_actual = f"${val_act:,.2f}"
        txt_meta = f"${val_meta:,.2f}"
    else:
        txt_actual = f"{val_act:,}"
        txt_meta = f"{val_meta:,}"

    delta = ((val_act - val_ant) / val_ant * 100) if val_ant > 0 else 0.0
    delta_class = "delta-pos" if delta >= 0 else "delta-neg"
    delta_sign = "+" if delta >= 0 else ""
    delta_icon = "▲" if delta >= 0 else "▼"

    cumplimiento = (val_act / val_meta * 100) if val_meta > 0 else 0.0
    ancho_barra = min(100.0, max(0.0, cumplimiento))

    return (
        f'<div class="kpi-card">'
        f'<div class="kpi-top">'
        f'<span class="kpi-label">{titulo}</span>'
        f'<span class="kpi-badge {badge_class}">{badge_txt}</span>'
        f'</div>'
        f'<div class="kpi-value">{txt_actual}</div>'
        f'<div class="kpi-progress-bg">'
        f'<div class="kpi-progress-fill {fill_class}" style="width: {ancho_barra:.1f}%;"></div>'
        f'</div>'
        f'<div class="kpi-bottom">'
        f'<span class="kpi-meta-text">🎯 Meta: <b>{txt_meta}</b> ({cumplimiento:.0f}%)</span>'
        f'<span class="kpi-delta-pill {delta_class}">{delta_icon} {delta_sign}{delta:.1f}% vs ant.</span>'
        f'</div>'
        f'</div>'
    )

def estilizar_figura(fig, titulo=""):
    """Formato común de alto contraste para visibilidad óptima."""
    fig.update_layout(
        template="plotly_white",
        title=dict(
            text=f"<b>{titulo}</b>",
            font=dict(size=14, color="#1A1D20", family="Plus Jakarta Sans"),
            x=0.02,
            y=0.96
        ),
        paper_bgcolor="#FFFFFF",
        plot_bgcolor="#FFFFFF",
        margin=dict(l=25, r=25, t=55, b=25),
        font=dict(family="Plus Jakarta Sans", color="#1A1D20", size=11),
        legend=dict(
            orientation="h",
            yanchor="bottom",
            y=1.02,
            xanchor="right",
            x=1,
            font=dict(size=10, color="#1A1D20")
        )
    )
    fig.update_xaxes(
        tickfont=dict(color="#1A1D20", size=10, family="Plus Jakarta Sans"),
        title_font=dict(color="#1A1D20", size=11, family="Plus Jakarta Sans"),
        gridcolor="#F0F2F5",
        zeroline=False,
        showline=True,
        linecolor="#E2E8F0"
    )
    fig.update_yaxes(
        tickfont=dict(color="#1A1D20", size=10, family="Plus Jakarta Sans"),
        title_font=dict(color="#1A1D20", size=11, family="Plus Jakarta Sans"),
        gridcolor="#F0F2F5",
        zeroline=False,
        showline=True,
        linecolor="#E2E8F0"
    )
    return fig

# --- 3. ESTILOS CSS INYECTADOS ---
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap');

    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }
    .stApp {
        background-color: #F5F7FA;
    }
    .block-container {
        padding-top: 1.8rem;
        padding-bottom: 2.5rem;
    }
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}

    header[data-testid="stHeader"] {
        background: transparent !important;
    }
    button[data-testid="stSidebarCollapsedControl"],
    div[data-testid="stSidebarCollapsedControl"] button {
        background-color: #FFFFFF !important;
        color: #1A1D20 !important;
        border: 1px solid #ECEFF2 !important;
        border-radius: 12px !important;
        box-shadow: 0 2px 10px rgba(0, 0, 0, 0.08) !important;
        margin-left: 12px !important;
        margin-top: 10px !important;
        visibility: visible !important;
        display: flex !important;
        align-items: center;
        justify-content: center;
        transition: all 0.2s ease !important;
    }
    button[data-testid="stSidebarCollapsedControl"]:hover,
    div[data-testid="stSidebarCollapsedControl"] button:hover {
        background-color: #FFB800 !important;
        color: #FFFFFF !important;
    }

    .dashboard-header {
        background: #FFFFFF;
        padding: 22px 28px;
        border-radius: 20px;
        display: flex;
        align-items: center;
        justify-content: space-between;
        margin-bottom: 24px;
        border: 1px solid #ECEFF2;
        box-shadow: 0px 4px 20px rgba(0, 0, 0, 0.03);
    }
    .dashboard-title-box h1 {
        font-size: 26px;
        font-weight: 800;
        color: #1A1D20 !important;
        margin: 0;
        letter-spacing: -0.5px;
    }
    .dashboard-title-box span {
        color: #FFB800 !important;
    }
    .dashboard-subtitle {
        color: #7A828A !important;
        font-size: 13px;
        font-weight: 500;
        margin-top: 4px;
    }

    .kpi-card {
        background: #FFFFFF;
        padding: 20px 22px;
        border-radius: 20px;
        border: 1px solid #ECEFF2;
        box-shadow: 0px 4px 18px rgba(0, 0, 0, 0.03);
        display: flex;
        flex-direction: column;
        justify-content: space-between;
        margin-bottom: 15px;
        min-height: 148px;
        transition: transform 0.2s ease, box-shadow 0.2s ease;
    }
    .kpi-card:hover {
        transform: translateY(-2px);
        box-shadow: 0px 8px 24px rgba(0, 0, 0, 0.06);
    }
    .kpi-top {
        display: flex;
        align-items: center;
        justify-content: space-between;
        margin-bottom: 4px;
    }
    .kpi-label {
        font-size: 11.5px;
        font-weight: 700;
        color: #8C94A0 !important;
        text-transform: uppercase;
        letter-spacing: 0.6px;
    }
    .kpi-value {
        font-size: 27px;
        font-weight: 800;
        color: #1A1D20 !important;
        margin: 2px 0 6px 0;
        line-height: 1.2;
    }
    .kpi-badge {
        padding: 5px 12px;
        border-radius: 20px;
        font-size: 11.5px;
        font-weight: 700;
    }
    .badge-yellow { background: #FFB800; color: #FFFFFF; }
    .badge-dark { background: #1E2229; color: #FFFFFF; }
    .badge-orange { background: #FF7A00; color: #FFFFFF; }
    .badge-purple { background: #6C5CE7; color: #FFFFFF; }

    .kpi-progress-bg {
        width: 100%;
        height: 6px;
        background-color: #F1F3F5;
        border-radius: 8px;
        overflow: hidden;
        margin: 6px 0 10px 0;
    }
    .kpi-progress-fill {
        height: 100%;
        border-radius: 8px;
        transition: width 0.3s ease;
    }
    .fill-yellow { background: linear-gradient(90deg, #FFB800, #FFA000); }
    .fill-dark { background: linear-gradient(90deg, #1E2229, #4B5563); }
    .fill-orange { background: linear-gradient(90deg, #FF7A00, #FF5500); }
    .fill-purple { background: linear-gradient(90deg, #6C5CE7, #805AD5); }

    .kpi-bottom {
        display: flex;
        align-items: center;
        justify-content: space-between;
        padding-top: 8px;
        border-top: 1px solid #F1F3F5;
        font-size: 11.5px;
    }
    .kpi-meta-text {
        color: #64748B !important;
        font-weight: 500;
    }
    .kpi-meta-text b {
        color: #1A1D20 !important;
        font-weight: 700;
    }
    .kpi-delta-pill {
        padding: 3px 9px;
        border-radius: 12px;
        font-size: 11px;
        font-weight: 700;
        display: inline-flex;
        align-items: center;
        white-space: nowrap;
    }
    .delta-pos { background: #E6F9F0; color: #00B894; }
    .delta-neg { background: #FEECEC; color: #E17055; }

    /* Tarjetas de Gráficos y Tablas */
    div[data-testid="stVerticalBlockBorderWrapper"] {
        background-color: #FFFFFF !important;
        border-radius: 22px !important;
        border: 1px solid #ECEFF2 !important;
        box-shadow: 0px 4px 18px rgba(0, 0, 0, 0.03) !important;
        padding: 16px !important;
    }

    button[data-baseweb="tab"] {
        font-weight: 700;
        color: #7A828A;
        font-size: 13.5px;
    }
    button[aria-selected="true"] {
        color: #1A1D20 !important;
        border-bottom-color: #FFB800 !important;
    }

    div.stButton > button:first-child {
        background: #FFB800;
        color: #FFFFFF;
        font-weight: 700;
        border: none;
        border-radius: 14px;
        padding: 10px 24px;
        transition: all 0.2s ease;
    }
    div.stButton > button:first-child:hover {
        background: #E5A600;
        color: #FFFFFF;
        transform: translateY(-1px);
        box-shadow: 0 4px 12px rgba(255, 184, 0, 0.35);
    }
</style>
""", unsafe_allow_html=True)

# --- 4. CONEXIÓN Y SERVICIOS CON GOOGLE DRIVE ---
SCOPES = ['https://www.googleapis.com/auth/drive']

@st.cache_resource
def get_drive_service():
    creds_dict = dict(st.secrets["gcp_service_account"])
    creds = service_account.Credentials.from_service_account_info(creds_dict, scopes=SCOPES)
    return build('drive', 'v3', credentials=creds)

@st.cache_data(ttl=60)
def load_data_from_drive(file_id, sheet_name=0):
    service = get_drive_service()
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    fh.seek(0)
    
    try:
        df = pd.read_excel(fh, sheet_name=sheet_name, engine='openpyxl')
    except Exception:
        fh.seek(0)
        df = pd.read_excel(fh, sheet_name=0, engine='openpyxl')
        
    df.columns = [str(c).strip().lower() for c in df.columns]
    return df

def save_ventas_to_drive(file_id, df_to_save):
    service = get_drive_service()
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    fh.seek(0)
    
    try:
        wb = openpyxl.load_workbook(fh)
    except Exception:
        wb = openpyxl.Workbook()

    ws = wb["Ventas"] if "Ventas" in wb.sheetnames else wb.active
    ws.title = "Ventas"
    ws.delete_rows(1, ws.max_row + 10)

    cols = ["id_transaccion", "fecha", "cliente", "ciudad", "categoria", "producto", "cantidad", "precio_unitario", "total_venta", "estado"]
    ws.append(cols)

    header_fill = PatternFill(start_color="1E2229", end_color="1E2229", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFB800")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    for row_idx, (_, row) in enumerate(df_to_save.iterrows(), start=2):
        cant = int(clean_val(row.get('cantidad', 0)))
        precio = float(clean_val(row.get('precio_unitario', 0.0)))
        formula_total = f"=G{row_idx}*H{row_idx}"
        
        row_values = [
            str(row.get('id_transaccion', f"TRX-{row_idx-1:04d}")),
            str(row.get('fecha', '')),
            str(row.get('cliente', '')),
            str(row.get('ciudad', '')),
            str(row.get('categoria', '')),
            str(row.get('producto', '')),
            cant,
            precio,
            formula_total,
            str(row.get('estado', 'Completado'))
        ]
        ws.append(row_values)
        ws.cell(row=row_idx, column=7).number_format = '#,##0'
        ws.cell(row=row_idx, column=8).number_format = '$#,##0.00'
        ws.cell(row=row_idx, column=9).number_format = '$#,##0.00'

        for c_i in range(1, len(cols) + 1):
            c_cell = ws.cell(row=row_idx, column=c_i)
            c_cell.border = thin_border
            if c_i in (1, 2, 10):
                c_cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)
    media = MediaIoBaseUpload(out_buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', resumable=True)
    service.files().update(fileId=file_id, media_body=media, supportsAllDrives=True).execute()

def save_inventario_to_drive(file_id, df_to_save):
    service = get_drive_service()
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    fh.seek(0)
    
    try:
        wb = openpyxl.load_workbook(fh)
    except Exception:
        wb = openpyxl.Workbook()

    ws = wb["Inventario"] if "Inventario" in wb.sheetnames else wb.active
    ws.title = "Inventario"
    ws.delete_rows(1, ws.max_row + 10)

    cols = ["id_inventario", "sku", "producto", "categoria", "almacen", "ciudad", "stock_actual", "stock_minimo", "stock_maximo", "costo_unitario", "valor_inventario", "estado_stock", "proveedor", "fecha_ultima_reposicion"]
    ws.append(cols)

    header_fill = PatternFill(start_color="1E2229", end_color="1E2229", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFB800")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    for row_idx, (_, row) in enumerate(df_to_save.iterrows(), start=2):
        s_act = int(clean_val(row.get('stock_actual', 0)))
        s_min = int(clean_val(row.get('stock_minimo', 0)))
        s_max = int(clean_val(row.get('stock_maximo', 0)))
        costo = float(clean_val(row.get('costo_unitario', 0.0)))
        formula_valor = f"=G{row_idx}*J{row_idx}"
        
        row_values = [
            str(row.get('id_inventario', f"INV-{row_idx-1:04d}")),
            str(row.get('sku', '')),
            str(row.get('producto', '')),
            str(row.get('categoria', '')),
            str(row.get('almacen', '')),
            str(row.get('ciudad', '')),
            s_act,
            s_min,
            s_max,
            costo,
            formula_valor,
            str(row.get('estado_stock', 'Óptimo')),
            str(row.get('proveedor', '')),
            str(row.get('fecha_ultima_reposicion', ''))
        ]
        ws.append(row_values)
        ws.cell(row=row_idx, column=7).number_format = '#,##0'
        ws.cell(row=row_idx, column=8).number_format = '#,##0'
        ws.cell(row=row_idx, column=9).number_format = '#,##0'
        ws.cell(row=row_idx, column=10).number_format = '$#,##0.00'
        ws.cell(row=row_idx, column=11).number_format = '$#,##0.00'

        for c_i in range(1, len(cols) + 1):
            c_cell = ws.cell(row=row_idx, column=c_i)
            c_cell.border = thin_border
            if c_i in (1, 2, 4, 6, 12, 14):
                c_cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)
    media = MediaIoBaseUpload(out_buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', resumable=True)
    service.files().update(fileId=file_id, media_body=media, supportsAllDrives=True).execute()

# --- 5. LECTURA DE IDENTIFICADORES DE DRIVE ---
FILE_ID_VENTAS = st.secrets["drive_settings"].get("file_id_ventas", st.secrets["drive_settings"].get("file_id"))
FILE_ID_INVENTARIO = st.secrets["drive_settings"].get("file_id_inventario", None)

# --- 6. BARRA LATERAL: SELECTOR DE MÓDULOS ---
st.sidebar.markdown("<h2 style='color:#FFFFFF; font-weight:800; margin-bottom: 2px;'>GLOBALTECH BI</h2>", unsafe_allow_html=True)
st.sidebar.caption("Suite Ejecutiva de Business Intelligence")

modulo_activo = st.sidebar.radio(
    "Selecciona el Módulo Operativo:",
    ["📈 Módulo Comercial (Ventas)", "📦 Módulo de Inventario & Almacenes"],
    index=0
)
st.sidebar.markdown("---")

# ==============================================================================
# --- CASO 1: MÓDULO COMERCIAL (VENTAS) ---
# ==============================================================================
if modulo_activo == "📈 Módulo Comercial (Ventas)":
    try:
        with st.spinner("Cargando base de Ventas desde Google Drive..."):
            df = load_data_from_drive(FILE_ID_VENTAS, sheet_name="Ventas")
    except Exception as e:
        st.error(f"Error al conectar con Google Drive (Ventas): {e}")
        st.stop()

    if 'fecha' in df.columns:
        df['fecha'] = pd.to_datetime(df['fecha'], errors='coerce').dt.strftime('%Y-%m-%d')
    df['cantidad'] = df['cantidad'].apply(clean_val).astype(int) if 'cantidad' in df.columns else 0
    df['precio_unitario'] = df['precio_unitario'].apply(clean_val) if 'precio_unitario' in df.columns else 0.0
    df['total_venta'] = df['total_venta'].apply(clean_val) if 'total_venta' in df.columns else 0.0
    mask = (df['total_venta'] == 0)
    df.loc[mask, 'total_venta'] = df.loc[mask, 'cantidad'] * df.loc[mask, 'precio_unitario']

    # Banner Superior
    st.markdown("""
    <div class="dashboard-header">
        <div class="dashboard-title-box">
            <h1>COMMERCIAL <span>Dashboard.</span></h1>
            <div class="dashboard-subtitle">GlobalTech BI • Ventas, Clientes y Facturación Estratégica</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # --- PESTAÑAS: GENERAL vs ANÁLISIS DETALLADO vs EDITOR ---
    tab_dash, tab_analisis, tab_edit = st.tabs([
        "📊 Dashboard General (KPIs)",
        "🔍 Análisis Detallado de Ventas",
        "✏️ Editor de Base de Datos"
    ])

    # FILTROS EN SIDEBAR (Comunes para ambas vistas de análisis)
    st.sidebar.markdown("<h3 style='color:#FFFFFF; font-weight:800;'>Filtros de Ventas</h3>", unsafe_allow_html=True)
    ciudades = ["Todas"] + sorted([c for c in df["ciudad"].dropna().unique().tolist() if str(c).strip()]) if "ciudad" in df.columns else ["Todas"]
    ciudad_sel = st.sidebar.selectbox("Ciudad:", ciudades)

    categorias = ["Todas"] + sorted([c for c in df["categoria"].dropna().unique().tolist() if str(c).strip()]) if "categoria" in df.columns else ["Todas"]
    cat_sel = st.sidebar.selectbox("Categoría:", categorias)

    estados = ["Todos"] + sorted([c for c in df["estado"].dropna().unique().tolist() if str(c).strip()]) if "estado" in df.columns else ["Todos"]
    estado_sel = st.sidebar.selectbox("Estado de Orden:", estados)

    with st.sidebar.expander("🎯 Metas Comerciales"):
        factor_meta = st.slider("Crecimiento Objetivo vs Período Anterior:", 5, 50, 15, 5, format="%d%%") / 100

    if st.sidebar.button("🔄 Actualizar Ventas", help="Refresca los datos en memoria"):
        st.cache_data.clear()
        st.rerun()

    # DataFrame Filtrado
    df_f = df.copy()
    if ciudad_sel != "Todas" and "ciudad" in df_f.columns:
        df_f = df_f[df_f["ciudad"] == ciudad_sel]
    if cat_sel != "Todas" and "categoria" in df_f.columns:
        df_f = df_f[df_f["categoria"] == cat_sel]
    if estado_sel != "Todos" and "estado" in df_f.columns:
        df_f = df_f[df_f["estado"] == estado_sel]

    # Cálculos globales
    total_ventas = float(df_f["total_venta"].sum())
    total_unidades = int(df_f["cantidad"].sum())
    num_ordenes = len(df_f)
    ticket_medio = (total_ventas / num_ordenes) if num_ordenes > 0 else 0.0

    # --------------------------------------------------------------------------
    # --- PESTAÑA 1: DASHBOARD GENERAL (KPIs y Matriz 3x2) ---
    # --------------------------------------------------------------------------
    with tab_dash:
        try:
            df_fechas = df_f.dropna(subset=['fecha']).copy()
            df_fechas['dt'] = pd.to_datetime(df_fechas['fecha'], errors='coerce')
            df_fechas = df_fechas.dropna(subset=['dt']).sort_values('dt')
            if len(df_fechas) >= 4:
                mid = len(df_fechas) // 2
                df_ant = df_fechas.iloc[:mid]
                ventas_ant = float(df_ant['total_venta'].sum())
                unidades_ant = int(df_ant['cantidad'].sum())
                ordenes_ant = len(df_ant)
                ticket_ant = (ventas_ant / ordenes_ant) if ordenes_ant > 0 else 0.0
            else:
                ventas_ant = total_ventas * 0.90
                unidades_ant = int(total_unidades * 0.90)
                ticket_ant = ticket_medio * 0.95
                ordenes_ant = int(num_ordenes * 0.92)
        except Exception:
            ventas_ant = total_ventas * 0.90
            unidades_ant = int(total_unidades * 0.90)
            ticket_ant = ticket_medio * 0.95
            ordenes_ant = int(num_ordenes * 0.92)

        meta_ventas = ventas_ant * (1 + factor_meta) if ventas_ant > 0 else total_ventas * 1.15
        meta_unidades = int(unidades_ant * (1 + factor_meta)) if unidades_ant > 0 else int(total_unidades * 1.15)
        meta_ticket = ticket_ant * (1 + (factor_meta * 0.6)) if ticket_ant > 0 else ticket_medio * 1.08
        meta_ordenes = int(ordenes_ant * (1 + factor_meta)) if ordenes_ant > 0 else int(num_ordenes * 1.15)

        k1, k2, k3, k4 = st.columns(4)
        with k1:
            st.markdown(render_kpi_card("Ingresos Totales", total_ventas, ventas_ant, meta_ventas, "Ventas", "badge-yellow", "fill-yellow", es_moneda=True), unsafe_allow_html=True)
        with k2:
            st.markdown(render_kpi_card("Unidades Vendidas", total_unidades, unidades_ant, meta_unidades, "Items", "badge-dark", "fill-dark", es_moneda=False), unsafe_allow_html=True)
        with k3:
            st.markdown(render_kpi_card("Ticket Promedio", ticket_medio, ticket_ant, meta_ticket, "Promedio", "badge-orange", "fill-orange", es_moneda=True), unsafe_allow_html=True)
        with k4:
            st.markdown(render_kpi_card("Total Órdenes", num_ordenes, ordenes_ant, meta_ordenes, "Órdenes", "badge-purple", "fill-purple", es_moneda=False), unsafe_allow_html=True)

        st.write("")

        # MATRIZ 3x2 DE VENTAS
        r1_c1, r1_c2 = st.columns(2)
        with r1_c1:
            with st.container(border=True):
                if len(df_f) > 0:
                    fig_disp = px.scatter(
                        df_f, x="cantidad", y="precio_unitario", size="total_venta",
                        color="categoria" if "categoria" in df_f.columns else None,
                        hover_name="producto" if "producto" in df_f.columns else None,
                        hover_data=["cliente", "total_venta"] if "cliente" in df_f.columns else ["total_venta"],
                        color_discrete_sequence=PALETA_COLORES, size_max=32,
                        labels={"cantidad": "Unidades por Orden", "precio_unitario": "Precio Unitario ($)", "categoria": "Categoría", "total_venta": "Monto Total ($)"}
                    )
                    fig_disp = estilizar_figura(fig_disp, "1. Dispersión Multidimensional 4D (Volumen vs Precio)")
                    st.plotly_chart(fig_disp, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        with r1_c2:
            with st.container(border=True):
                if len(df_f) > 0 and all(c in df_f.columns for c in ["ciudad", "categoria", "producto"]):
                    df_tree = df_f.copy()
                    for c in ["ciudad", "categoria", "producto"]:
                        df_tree[c] = df_tree[c].fillna("S/D").astype(str)
                    fig_tree = px.treemap(
                        df_tree, path=["ciudad", "categoria", "producto"], values="total_venta",
                        color="categoria", color_discrete_sequence=PALETA_COLORES
                    )
                    fig_tree = estilizar_figura(fig_tree, "2. Treemap Jerárquico (Ciudad ➔ Categoría ➔ Producto)")
                    fig_tree.update_traces(textinfo="label+value+percent entry", texttemplate="<b>%{label}</b><br>$%{value:,.0f}")
                    st.plotly_chart(fig_tree, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        r2_c1, r2_c2 = st.columns(2)
        with r2_c1:
            with st.container(border=True):
                if "producto" in df_f.columns and len(df_f) > 0:
                    v_prod = df_f.groupby("producto", as_index=False)["total_venta"].sum().sort_values("total_venta", ascending=False).head(10)
                    v_prod["cum_sum"] = v_prod["total_venta"].cumsum()
                    v_prod["cum_pct"] = (v_prod["cum_sum"] / v_prod["total_venta"].sum()) * 100

                    fig_pareto = make_subplots(specs=[[{"secondary_y": True}]])
                    fig_pareto.add_trace(go.Bar(x=v_prod["producto"], y=v_prod["total_venta"], name="Ventas ($)", marker_color="#FFB800", text=v_prod["total_venta"], texttemplate="$%{text:,.0f}", textposition="outside"), secondary_y=False)
                    fig_pareto.add_trace(go.Scatter(x=v_prod["producto"], y=v_prod["cum_pct"], name="% Acumulado", mode="lines+markers", line=dict(color="#1E2229", width=3), marker=dict(size=6, color="#FF7A00")), secondary_y=True)
                    fig_pareto.add_hline(y=80, line_dash="dot", line_color="#E17055", line_width=2, annotation_text="Regla 80%", secondary_y=True)
                    fig_pareto = estilizar_figura(fig_pareto, "3. Diagrama de Pareto 80/20 (Top Productos)")
                    fig_pareto.update_yaxes(title_text="Ventas ($)", secondary_y=False, showgrid=True, gridcolor="#F0F2F5")
                    fig_pareto.update_yaxes(title_text="% Acumulado", range=[0, 110], secondary_y=True, showgrid=False)
                    fig_pareto.update_xaxes(tickangle=-25)
                    st.plotly_chart(fig_pareto, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        with r2_c2:
            with st.container(border=True):
                if "fecha" in df_f.columns and len(df_f) > 0:
                    v_tiempo = df_f.groupby("fecha", as_index=False)["total_venta"].sum().sort_values("fecha")
                    fig_line = px.line(v_tiempo, x="fecha", y="total_venta", text="total_venta", color_discrete_sequence=["#FF7A00"], markers=True, labels={"total_venta": "Ingresos ($)", "fecha": "Fecha"})
                    fig_line = estilizar_figura(fig_line, "4. Tendencia Histórica de Ventas")
                    fig_line.update_traces(line=dict(width=3, shape="spline"), marker=dict(size=7, color="#1E2229"), texttemplate='$%{text:,.0f}', textposition='top center')
                    st.plotly_chart(fig_line, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        r3_c1, r3_c2 = st.columns(2)
        with r3_c1:
            with st.container(border=True):
                if len(df_f) > 0 and all(c in df_f.columns for c in ["ciudad", "categoria", "estado"]):
                    ciudades_list = list(df_f["ciudad"].dropna().unique())
                    categorias_list = list(df_f["categoria"].dropna().unique())
                    estados_list = list(df_f["estado"].dropna().unique())
                    all_nodes = ciudades_list + categorias_list + estados_list
                    node_indices = {name: i for i, name in enumerate(all_nodes)}

                    df_l1 = df_f.groupby(["ciudad", "categoria"], as_index=False)["total_venta"].sum()
                    df_l2 = df_f.groupby(["categoria", "estado"], as_index=False)["total_venta"].sum()

                    srcs = [node_indices[c] for c in df_l1["ciudad"]] + [node_indices[cat] for cat in df_l2["categoria"]]
                    tgts = [node_indices[cat] for cat in df_l1["categoria"]] + [node_indices[est] for est in df_l2["estado"]]
                    vals = list(df_l1["total_venta"]) + list(df_l2["total_venta"])

                    node_colors = ["#FFB800"] * len(ciudades_list) + ["#1E2229"] * len(categorias_list) + ["#6C5CE7"] * len(estados_list)
                    fig_sankey = go.Figure(data=[go.Sankey(
                        node=dict(pad=15, thickness=18, line=dict(color="#ECEFF2", width=1), label=all_nodes, color=node_colors),
                        link=dict(source=srcs, target=tgts, value=vals, color="rgba(255, 184, 0, 0.28)")
                    )])
                    fig_sankey = estilizar_figura(fig_sankey, "5. Diagrama de Flujo Sankey (Ciudad ➔ Categoría ➔ Estado)")
                    st.plotly_chart(fig_sankey, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        with r3_c2:
            with st.container(border=True):
                if "categoria" in df_f.columns and len(df_f) > 0:
                    df_radar = df_f.groupby("categoria").agg(ventas=("total_venta", "sum"), unidades=("cantidad", "sum"), ordenes=("id_transaccion", "count"), ticket=("total_venta", "mean")).reset_index()
                    metricas = ["ventas", "unidades", "ordenes", "ticket"]
                    nombres_ejes = ["Facturación ($)", "Unidades", "Nº Órdenes", "Ticket Medio"]
                    for m in metricas:
                        max_v = df_radar[m].max()
                        df_radar[m + "_norm"] = (df_radar[m] / max_v * 100) if max_v > 0 else 0

                    fig_radar = go.Figure()
                    for idx, cat_name in enumerate(df_radar["categoria"].unique()):
                        row = df_radar[df_radar["categoria"] == cat_name].iloc[0]
                        valores_radar = [row[m + "_norm"] for m in metricas]
                        valores_radar.append(valores_radar[0])
                        ejes_cerrados = nombres_ejes + [nombres_ejes[0]]
                        fig_radar.add_trace(go.Scatterpolar(r=valores_radar, theta=ejes_cerrados, fill='toself', name=str(cat_name), line=dict(color=PALETA_COLORES[idx % len(PALETA_COLORES)], width=2.5), opacity=0.65))

                    fig_radar = estilizar_figura(fig_radar, "6. Radar Polar Multicriterio 360° por Categoría")
                    fig_radar.update_layout(polar=dict(radialaxis=dict(visible=True, range=[0, 105], tickfont=dict(size=9, color="#7A828A"), gridcolor="#F0F2F5"), bgcolor="#FFFFFF"))
                    st.plotly_chart(fig_radar, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

    # --------------------------------------------------------------------------
    # --- PESTAÑA 2: ANÁLISIS DETALLADO DE VENTAS (TABLAS Y PROFUNDIDAD) ---
    # --------------------------------------------------------------------------
    with tab_analisis:
        st.markdown("""
        <div style='background: #FFFFFF; padding: 18px 24px; border-radius: 18px; border: 1px solid #ECEFF2; margin-bottom: 20px; box-shadow: 0px 4px 18px rgba(0, 0, 0, 0.02);'>
            <h3 style='margin:0; color:#1A1D20; font-weight:800;'>🔍 Análisis Analítico de Rendimiento y Clientes</h3>
            <p style='color:#7A828A; font-size:13px; margin-top:4px;'>Profundiza en la cartera de clientes, concentración geográfica cruzada y dispersión de precios.</p>
        </div>
        """, unsafe_allow_html=True)

        if len(df_f) > 0:
            # 1. Mini-KPIs de hallazgos
            top_cli_nom = df_f.groupby("cliente")["total_venta"].sum().idxmax() if "cliente" in df_f.columns else "N/A"
            top_cli_val = df_f.groupby("cliente")["total_venta"].sum().max() if "cliente" in df_f.columns else 0.0

            top_prod_nom = df_f.groupby("producto")["total_venta"].sum().idxmax() if "producto" in df_f.columns else "N/A"
            top_prod_val = df_f.groupby("producto")["total_venta"].sum().max() if "producto" in df_f.columns else 0.0

            top_ciu_nom = df_f.groupby("ciudad")["total_venta"].sum().idxmax() if "ciudad" in df_f.columns else "N/A"
            top_ciu_val = df_f.groupby("ciudad")["total_venta"].sum().max() if "ciudad" in df_f.columns else 0.0

            prom_unidades_orden = df_f["cantidad"].mean()

            ak1, ak2, ak3, ak4 = st.columns(4)
            with ak1:
                st.markdown(f"""
                <div class="kpi-card" style="min-height: 110px;">
                    <span class="kpi-label">🏆 Cliente N° 1</span>
                    <span style="font-size:18px; font-weight:800; color:#1A1D20; margin-top:4px;">{top_cli_nom}</span>
                    <span style="font-size:12px; color:#FFB800; font-weight:700;">${top_cli_val:,.2f}</span>
                </div>
                """, unsafe_allow_html=True)
            with ak2:
                st.markdown(f"""
                <div class="kpi-card" style="min-height: 110px;">
                    <span class="kpi-label">⭐ Producto Estrella</span>
                    <span style="font-size:18px; font-weight:800; color:#1A1D20; margin-top:4px;">{top_prod_nom}</span>
                    <span style="font-size:12px; color:#FF7A00; font-weight:700;">${top_prod_val:,.2f}</span>
                </div>
                """, unsafe_allow_html=True)
            with ak3:
                st.markdown(f"""
                <div class="kpi-card" style="min-height: 110px;">
                    <span class="kpi-label">📍 Plaza Principal</span>
                    <span style="font-size:18px; font-weight:800; color:#1A1D20; margin-top:4px;">{top_ciu_nom}</span>
                    <span style="font-size:12px; color:#1E2229; font-weight:700;">${top_ciu_val:,.2f}</span>
                </div>
                """, unsafe_allow_html=True)
            with ak4:
                st.markdown(f"""
                <div class="kpi-card" style="min-height: 110px;">
                    <span class="kpi-label">📦 Unidades / Orden</span>
                    <span style="font-size:22px; font-weight:800; color:#1A1D20; margin-top:4px;">{prom_unidades_orden:.1f}</span>
                    <span style="font-size:12px; color:#6C5CE7; font-weight:700;">Promedio físico</span>
                </div>
                """, unsafe_allow_html=True)

            st.write("")

            # 2. FILA DE CLIENTES: GRÁFICO RANKING + TABLA DINÁMICA
            ac1, ac2 = st.columns([1, 1.2])

            with ac1:
                with st.container(border=True):
                    # Gráfico Top 10 Clientes
                    if "cliente" in df_f.columns:
                        v_cli = df_f.groupby("cliente", as_index=False)["total_venta"].sum().sort_values("total_venta", ascending=True).tail(10)
                        fig_cli = px.bar(
                            v_cli,
                            x="total_venta",
                            y="cliente",
                            orientation="h",
                            text="total_venta",
                            color_discrete_sequence=["#1E2229"],
                            labels={"total_venta": "Facturación ($)", "cliente": ""}
                        )
                        fig_cli = estilizar_figura(fig_cli, "Ranking: Top 10 Clientes por Facturación")
                        fig_cli.update_traces(
                            texttemplate='$%{text:,.0f}',
                            textposition='outside',
                            textfont=dict(color="#1A1D20", size=10, family="Plus Jakarta Sans")
                        )
                        max_cli = v_cli["total_venta"].max()
                        fig_cli.update_xaxes(range=[0, max_cli * 1.30])
                        st.plotly_chart(fig_cli, width="stretch", theme=None)

            with ac2:
                with st.container(border=True):
                    # Tabla Resumen de Clientes
                    st.markdown("<b style='color:#1A1D20; font-size:14px;'>📋 Matriz de Desempeño por Cliente</b>", unsafe_allow_html=True)
                    st.caption("Detalle de volumen, ticket promedio y porcentaje de aportación al negocio.")
                    
                    if "cliente" in df_f.columns:
                        tabla_cli = df_f.groupby("cliente").agg(
                            Ordenes=("id_transaccion", "count"),
                            Unidades=("cantidad", "sum"),
                            Total_Venta=("total_venta", "sum"),
                            Ticket_Medio=("total_venta", "mean")
                        ).reset_index()

                        # % sobre el total filtrado
                        sum_tot = tabla_cli["Total_Venta"].sum()
                        tabla_cli["Participacion"] = (tabla_cli["Total_Venta"] / sum_tot) * 100 if sum_tot > 0 else 0
                        tabla_cli = tabla_cli.sort_values("Total_Venta", ascending=False)

                        st.dataframe(
                            tabla_cli,
                            column_config={
                                "cliente": st.column_config.TextColumn("Cliente"),
                                "Ordenes": st.column_config.NumberColumn("Órdenes", format="%d"),
                                "Unidades": st.column_config.NumberColumn("Unidades", format="%d"),
                                "Total_Venta": st.column_config.NumberColumn("Total Facturado", format="$ %,.2f"),
                                "Ticket_Medio": st.column_config.NumberColumn("Ticket Medio", format="$ %,.2f"),
                                "Participacion": st.column_config.ProgressColumn("Participación %", format="%.1f%%", min_value=0, max_value=100)
                            },
                            hide_index=True,
                            width="stretch",
                            height=380
                        )

            # 3. FILA DE MATRIZ CRUZADA (HEATMAP) + DISPERSIÓN DE PRECIOS (BOXPLOT)
            st.write("")
            ah1, ah2 = st.columns(2)

            with ah1:
                with st.container(border=True):
                    # Heatmap Cruzado Ciudad vs Categoría
                    if "ciudad" in df_f.columns and "categoria" in df_f.columns:
                        pivot_cc = df_f.pivot_table(index="ciudad", columns="categoria", values="total_venta", aggfunc="sum", fill_value=0)
                        fig_heat = px.imshow(
                            pivot_cc,
                            text_auto="$,.0f",
                            aspect="auto",
                            color_continuous_scale=[[0, "#FFFFFF"], [0.4, "#FFF3CD"], [0.8, "#FFB800"], [1, "#1E2229"]],
                            labels=dict(x="Categoría", y="Ciudad", color="Facturación ($)")
                        )
                        fig_heat = estilizar_figura(fig_heat, "Mapa de Calor: Facturación Cruzada (Ciudad vs Categoría)")
                        fig_heat.update_xaxes(side="bottom")
                        st.plotly_chart(fig_heat, width="stretch", theme=None)
                    else:
                        st.info("Sin columnas para generar el mapa de calor.")

            with ah2:
                with st.container(border=True):
                    # Boxplot de dispersión de precios unitarios
                    if "categoria" in df_f.columns and "precio_unitario" in df_f.columns:
                        fig_box = px.box(
                            df_f,
                            x="categoria",
                            y="precio_unitario",
                            color="categoria",
                            points="all",
                            color_discrete_sequence=PALETA_COLORES,
                            labels={"precio_unitario": "Precio Unitario ($)", "categoria": "Categoría"}
                        )
                        fig_box = estilizar_figura(fig_box, "Dispersión de Precios Unitarios por Categoría (Boxplot)")
                        fig_box.update_layout(showlegend=False)
                        st.plotly_chart(fig_box, width="stretch", theme=None)
                    else:
                        st.info("Sin registros de precios.")

            # 4. TABLA AUDITABLE DETALLE LÍNEA POR LÍNEA
            st.write("")
            with st.container(border=True):
                st.markdown("<b style='color:#1A1D20; font-size:15px;'>🔎 Explorador Maestro de Transacciones Filtradas</b>", unsafe_allow_html=True)
                st.caption(f"Mostrando {len(df_f)} transacciones que coinciden con los filtros del panel lateral.")
                
                cols_mostrar = [c for c in ["id_transaccion", "fecha", "cliente", "ciudad", "categoria", "producto", "cantidad", "precio_unitario", "total_venta", "estado"] if c in df_f.columns]
                st.dataframe(
                    df_f[cols_mostrar],
                    column_config={
                        "precio_unitario": st.column_config.NumberColumn("Precio ($)", format="$ %,.2f"),
                        "total_venta": st.column_config.NumberColumn("Total ($)", format="$ %,.2f"),
                        "cantidad": st.column_config.NumberColumn("Cantidad", format="%d"),
                        "fecha": st.column_config.DateColumn("Fecha", format="YYYY-MM-DD")
                    },
                    hide_index=True,
                    width="stretch",
                    height=320
                )
        else:
            st.info("No hay datos disponibles para los filtros seleccionados.")

    # --------------------------------------------------------------------------
    # --- PESTAÑA 3: EDITOR DE BASE DE DATOS ---
    # --------------------------------------------------------------------------
    with tab_edit:
        st.markdown("<h3 style='margin:0; color:#1A1D20; font-weight:800;'>Editor Directo: Ventas</h3><p style='color:#7A828A; font-size:13px;'>Modifica valores en la tabla. Las fórmulas se sincronizan en Google Drive.</p>", unsafe_allow_html=True)
        df_editado = st.data_editor(df, num_rows="dynamic", width="stretch", hide_index=True, disabled=["total_venta"])
        if st.button("💾 Guardar cambios de Ventas en Google Drive", type="primary"):
            try:
                with st.spinner("Guardando en Google Drive..."):
                    save_ventas_to_drive(FILE_ID_VENTAS, df_editado)
                st.success("¡Base de Ventas actualizada correctamente!")
                st.cache_data.clear()
                st.rerun()
            except Exception as e:
                st.error(f"Error al guardar: {e}")

# ==============================================================================
# --- CASO 2: MÓDULO DE INVENTARIO & ALMACENES ---
# ==============================================================================
else:
    if not FILE_ID_INVENTARIO:
        st.warning("⚠️ Falta configurar `file_id_inventario` en los Secrets de Streamlit Cloud.")
        st.info("1. Sube el archivo `base_datos_inventario.xlsx` a tu Google Drive.\n2. Abre Streamlit Cloud > Settings > Secrets y añade:\n\n```toml\n[drive_settings]\nfile_id_inventario = 'TU_ID_AQUÍ'\n```")
        st.stop()

    try:
        with st.spinner("Cargando base de Inventario desde Google Drive..."):
            df_inv = load_data_from_drive(FILE_ID_INVENTARIO, sheet_name="Inventario")
    except Exception as e:
        st.error(f"Error al conectar con Google Drive (Inventario): {e}")
        st.stop()

    df_inv['stock_actual'] = df_inv['stock_actual'].apply(clean_val).astype(int) if 'stock_actual' in df_inv.columns else 0
    df_inv['stock_minimo'] = df_inv['stock_minimo'].apply(clean_val).astype(int) if 'stock_minimo' in df_inv.columns else 0
    df_inv['stock_maximo'] = df_inv['stock_maximo'].apply(clean_val).astype(int) if 'stock_maximo' in df_inv.columns else 0
    df_inv['costo_unitario'] = df_inv['costo_unitario'].apply(clean_val) if 'costo_unitario' in df_inv.columns else 0.0
    df_inv['valor_inventario'] = df_inv['valor_inventario'].apply(clean_val) if 'valor_inventario' in df_inv.columns else 0.0
    mask_inv = (df_inv['valor_inventario'] == 0)
    df_inv.loc[mask_inv, 'valor_inventario'] = df_inv.loc[mask_inv, 'stock_actual'] * df_inv.loc[mask_inv, 'costo_unitario']

    st.markdown("""
    <div class="dashboard-header">
        <div class="dashboard-title-box">
            <h1>INVENTORY <span>Dashboard.</span></h1>
            <div class="dashboard-subtitle">GlobalTech BI • Gestión de Stock, Valorización y Cadena de Suministro</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    tab_inv_dash, tab_inv_edit = st.tabs(["📊 Dashboard de Inventario", "✏️ Editor de Base de Datos"])

    with tab_inv_dash:
        st.sidebar.markdown("<h3 style='color:#FFFFFF; font-weight:800;'>Filtros de Almacén</h3>", unsafe_allow_html=True)
        almacenes = ["Todos"] + sorted([a for a in df_inv["almacen"].dropna().unique().tolist() if str(a).strip()]) if "almacen" in df_inv.columns else ["Todos"]
        alm_sel = st.sidebar.selectbox("Almacén:", almacenes)

        categorias_inv = ["Todas"] + sorted([c for c in df_inv["categoria"].dropna().unique().tolist() if str(c).strip()]) if "categoria" in df_inv.columns else ["Todas"]
        cat_inv_sel = st.sidebar.selectbox("Categoría de Producto:", categorias_inv)

        estados_inv = ["Todos"] + sorted([e for e in df_inv["estado_stock"].dropna().unique().tolist() if str(e).strip()]) if "estado_stock" in df_inv.columns else ["Todos"]
        est_inv_sel = st.sidebar.selectbox("Estado del Stock:", estados_inv)

        with st.sidebar.expander("🎯 Capacidad Objetivo"):
            capacidad_almacen_kpi = st.slider("Límite de Inversión en Stock ($):", 100000, 500000, 250000, 25000, format="$%d")

        if st.sidebar.button("🔄 Actualizar Inventario", help="Refresca los datos en memoria"):
            st.cache_data.clear()
            st.rerun()

        df_inv_f = df_inv.copy()
        if alm_sel != "Todos" and "almacen" in df_inv_f.columns:
            df_inv_f = df_inv_f[df_inv_f["almacen"] == alm_sel]
        if cat_inv_sel != "Todas" and "categoria" in df_inv_f.columns:
            df_inv_f = df_inv_f[df_inv_f["categoria"] == cat_inv_sel]
        if est_inv_sel != "Todos" and "estado_stock" in df_inv_f.columns:
            df_inv_f = df_inv_f[df_inv_f["estado_stock"] == est_inv_sel]

        val_total_stock = float(df_inv_f["valor_inventario"].sum())
        total_items_stock = int(df_inv_f["stock_actual"].sum())
        items_criticos = int(len(df_inv_f[df_inv_f["estado_stock"].astype(str).str.contains("Crítico|Bajo", case=False, na=False)]))
        total_skus = int(len(df_inv_f))

        k1, k2, k3, k4 = st.columns(4)
        with k1:
            st.markdown(render_kpi_card("Capital Inmovilizado", val_total_stock, val_total_stock * 0.92, float(capacidad_almacen_kpi), "Valorización", "badge-yellow", "fill-yellow", es_moneda=True), unsafe_allow_html=True)
        with k2:
            st.markdown(render_kpi_card("Stock Total Físico", total_items_stock, int(total_items_stock * 0.95), int(total_items_stock * 1.10), "Unidades", "badge-dark", "fill-dark", es_moneda=False), unsafe_allow_html=True)
        with k3:
            st.markdown(render_kpi_card("Alertas de Reposición", items_criticos, int(items_criticos * 1.2), max(1, int(items_criticos * 0.5)), "Críticos", "badge-orange", "fill-orange", es_moneda=False), unsafe_allow_html=True)
        with k4:
            st.markdown(render_kpi_card("Líneas de Producto (SKUs)", total_skus, total_skus, total_skus, "Catálogo", "badge-purple", "fill-purple", es_moneda=False), unsafe_allow_html=True)

        st.write("")

        # MATRIZ 3x2 DE INVENTARIOS
        i1_c1, i1_c2 = st.columns(2)
        with i1_c1:
            with st.container(border=True):
                if len(df_inv_f) > 0:
                    fig_disp_inv = px.scatter(
                        df_inv_f, x="stock_actual", y="costo_unitario", size="valor_inventario",
                        color="categoria" if "categoria" in df_inv_f.columns else None,
                        hover_name="producto" if "producto" in df_inv_f.columns else None,
                        hover_data=["almacen", "estado_stock", "valor_inventario"] if "almacen" in df_inv_f.columns else ["valor_inventario"],
                        color_discrete_sequence=PALETA_COLORES, size_max=32,
                        labels={"stock_actual": "Stock Físico Disponible", "costo_unitario": "Costo de Adquisición ($)", "categoria": "Categoría", "valor_inventario": "Valorización ($)"}
                    )
                    fig_disp_inv = estilizar_figura(fig_disp_inv, "1. Dispersión 4D de Stock (Unidades vs Costo vs Valor Total)")
                    st.plotly_chart(fig_disp_inv, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        with i1_c2:
            with st.container(border=True):
                if len(df_inv_f) > 0 and all(c in df_inv_f.columns for c in ["almacen", "categoria", "producto"]):
                    df_tree_inv = df_inv_f.copy()
                    for c in ["almacen", "categoria", "producto"]:
                        df_tree_inv[c] = df_tree_inv[c].fillna("S/D").astype(str)
                    fig_tree_inv = px.treemap(
                        df_tree_inv, path=["almacen", "categoria", "producto"], values="valor_inventario",
                        color="categoria", color_discrete_sequence=PALETA_COLORES
                    )
                    fig_tree_inv = estilizar_figura(fig_tree_inv, "2. Treemap de Inventario (Almacén ➔ Categoría ➔ Producto)")
                    fig_tree_inv.update_traces(textinfo="label+value+percent entry", texttemplate="<b>%{label}</b><br>$%{value:,.0f}")
                    st.plotly_chart(fig_tree_inv, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        i2_c1, i2_c2 = st.columns(2)
        with i2_c1:
            with st.container(border=True):
                if "producto" in df_inv_f.columns and len(df_inv_f) > 0:
                    v_inv_prod = df_inv_f.groupby("producto", as_index=False)["valor_inventario"].sum().sort_values("valor_inventario", ascending=False).head(10)
                    v_inv_prod["cum_sum"] = v_inv_prod["valor_inventario"].cumsum()
                    v_inv_prod["cum_pct"] = (v_inv_prod["cum_sum"] / v_inv_prod["valor_inventario"].sum()) * 100

                    fig_pareto_inv = make_subplots(specs=[[{"secondary_y": True}]])
                    fig_pareto_inv.add_trace(go.Bar(x=v_inv_prod["producto"], y=v_inv_prod["valor_inventario"], name="Valorización ($)", marker_color="#FFB800", text=v_inv_prod["valor_inventario"], texttemplate="$%{text:,.0f}", textposition="outside"), secondary_y=False)
                    fig_pareto_inv.add_trace(go.Scatter(x=v_inv_prod["producto"], y=v_inv_prod["cum_pct"], name="% Acumulado", mode="lines+markers", line=dict(color="#1E2229", width=3), marker=dict(size=6, color="#FF7A00")), secondary_y=True)
                    fig_pareto_inv.add_hline(y=80, line_dash="dot", line_color="#E17055", line_width=2, annotation_text="Corte 80%", secondary_y=True)
                    fig_pareto_inv = estilizar_figura(fig_pareto_inv, "3. Pareto 80/20: SKUs con Mayor Capital Inmovilizado")
                    fig_pareto_inv.update_yaxes(title_text="Valorización ($)", secondary_y=False, showgrid=True, gridcolor="#F0F2F5")
                    fig_pareto_inv.update_yaxes(title_text="% Acumulado", range=[0, 110], secondary_y=True, showgrid=False)
                    fig_pareto_inv.update_xaxes(tickangle=-25)
                    st.plotly_chart(fig_pareto_inv, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        with i2_c2:
            with st.container(border=True):
                if "estado_stock" in df_inv_f.columns and len(df_inv_f) > 0:
                    v_est_inv = df_inv_f.groupby("estado_stock", as_index=False)["valor_inventario"].sum()
                    fig_donut_inv = px.pie(
                        v_est_inv, values="valor_inventario", names="estado_stock", hole=0.62,
                        color_discrete_sequence=["#00B894", "#E17055", "#FF7A00", "#FFB800"]
                    )
                    fig_donut_inv = estilizar_figura(fig_donut_inv, "4. Distribución de Salud del Stock")
                    fig_donut_inv.update_traces(textinfo="label+percent", textfont=dict(size=11, color="#FFFFFF"))
                    st.plotly_chart(fig_donut_inv, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        i3_c1, i3_c2 = st.columns(2)
        with i3_c1:
            with st.container(border=True):
                if len(df_inv_f) > 0 and all(c in df_inv_f.columns for c in ["almacen", "categoria", "estado_stock"]):
                    alms_list = list(df_inv_f["almacen"].dropna().unique())
                    cats_list = list(df_inv_f["categoria"].dropna().unique())
                    ests_list = list(df_inv_f["estado_stock"].dropna().unique())
                    all_inv_nodes = alms_list + cats_list + ests_list
                    inv_node_indices = {name: i for i, name in enumerate(all_inv_nodes)}

                    df_il1 = df_inv_f.groupby(["almacen", "categoria"], as_index=False)["valor_inventario"].sum()
                    df_il2 = df_inv_f.groupby(["categoria", "estado_stock"], as_index=False)["valor_inventario"].sum()

                    s_srcs = [inv_node_indices[a] for a in df_il1["almacen"]] + [inv_node_indices[c] for c in df_il2["categoria"]]
                    s_tgts = [inv_node_indices[c] for c in df_il1["categoria"]] + [inv_node_indices[e] for e in df_il2["estado_stock"]]
                    s_vals = list(df_il1["valor_inventario"]) + list(df_il2["valor_inventario"])

                    node_inv_colors = ["#FFB800"] * len(alms_list) + ["#1E2229"] * len(cats_list) + ["#00B894"] * len(ests_list)
                    fig_sankey_inv = go.Figure(data=[go.Sankey(
                        node=dict(pad=15, thickness=18, line=dict(color="#ECEFF2", width=1), label=all_inv_nodes, color=node_inv_colors),
                        link=dict(source=s_srcs, target=s_tgts, value=s_vals, color="rgba(255, 184, 0, 0.28)")
                    )])
                    fig_sankey_inv = estilizar_figura(fig_sankey_inv, "5. Flujo Sankey: Almacén ➔ Categoría ➔ Estado de Stock")
                    st.plotly_chart(fig_sankey_inv, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

        with i3_c2:
            with st.container(border=True):
                if "almacen" in df_inv_f.columns and len(df_inv_f) > 0:
                    df_radar_alm = df_inv_f.groupby("almacen").agg(
                        valor=("valor_inventario", "sum"), unidades=("stock_actual", "sum"),
                        skus=("sku", "count"), seguridad=("stock_minimo", "sum")
                    ).reset_index()

                    m_inv = ["valor", "unidades", "skus", "seguridad"]
                    ejes_inv = ["Valorización ($)", "Stock Físico", "Variedad SKUs", "Stock de Seguridad"]

                    for m in m_inv:
                        max_m = df_radar_alm[m].max()
                        df_radar_alm[m + "_norm"] = (df_radar_alm[m] / max_m * 100) if max_m > 0 else 0

                    fig_radar_inv = go.Figure()
                    for idx, alm_n in enumerate(df_radar_alm["almacen"].unique()):
                        row_a = df_radar_alm[df_radar_alm["almacen"] == alm_n].iloc[0]
                        v_rad = [row_a[m + "_norm"] for m in m_inv]
                        v_rad.append(v_rad[0])
                        ejes_c = ejes_inv + [ejes_inv[0]]
                        fig_radar_inv.add_trace(go.Scatterpolar(r=v_rad, theta=ejes_c, fill='toself', name=str(alm_n), line=dict(color=PALETA_COLORES[idx % len(PALETA_COLORES)], width=2.5), opacity=0.60))

                    fig_radar_inv = estilizar_figura(fig_radar_inv, "6. Radar Polar 360°: Rendimiento y Balance por Almacén")
                    fig_radar_inv.update_layout(polar=dict(radialaxis=dict(visible=True, range=[0, 105], tickfont=dict(size=9, color="#7A828A"), gridcolor="#F0F2F5"), bgcolor="#FFFFFF"))
                    st.plotly_chart(fig_radar_inv, width="stretch", theme=None)
                else:
                    st.info("Sin registros.")

    with tab_inv_edit:
        st.markdown("<h3 style='margin:0; color:#1A1D20; font-weight:800;'>Editor Directo: Inventario</h3><p style='color:#7A828A; font-size:13px;'>Edita cantidades, costos y proveedores. Los totales se calculan automáticamente con fórmulas en Google Drive.</p>", unsafe_allow_html=True)
        df_inv_editado = st.data_editor(df_inv, num_rows="dynamic", width="stretch", hide_index=True, disabled=["valor_inventario"])
        if st.button("💾 Guardar cambios de Inventario en Google Drive", type="primary"):
            try:
                with st.spinner("Guardando en Google Drive..."):
                    save_inventario_to_drive(FILE_ID_INVENTARIO, df_inv_editado)
                st.success("¡Base de Inventario actualizada en Google Drive!")
                st.cache_data.clear()
                st.rerun()
            except Exception as e:
                st.error(f"Error al guardar: {e}")
